# =========================================================
# excel_utils.py
# =========================================================

from openpyxl import load_workbook


class ExcelUtils:

    # =====================================================
    # GET COMPLETE SHEET DATA
    # =====================================================

    @staticmethod
    def get_all_data(
            path,
            sheet_name
    ):

        workbook = load_workbook(path)

        sheet = workbook[sheet_name]

        data = []

        for row in sheet.iter_rows(
                min_row=2,
                values_only=True
        ):

            data.append(row)

        return data

    # =====================================================
    # GET DATA USING TC_ID
    # =====================================================

    @staticmethod
    def get_data_by_tc_id(
            path,
            sheet_name,
            tc_id
    ):

        workbook = load_workbook(path)

        sheet = workbook[sheet_name]

        for row in sheet.iter_rows(
                min_row=2,
                values_only=True
        ):

            if row[0] == tc_id:

                return row

        return None