from openpyxl import load_workbook


class ExcelUtils:

    @staticmethod
    def get_row_count(path, sheet_name):

        workbook = load_workbook(path)

        sheet = workbook[sheet_name]

        return sheet.max_row

    @staticmethod
    def get_column_count(path, sheet_name):

        workbook = load_workbook(path)

        sheet = workbook[sheet_name]

        return sheet.max_column

    @staticmethod
    def read_data(path, sheet_name, row_num, column_num):

        workbook = load_workbook(path)

        sheet = workbook[sheet_name]

        return sheet.cell(
            row=row_num,
            column=column_num
        ).value

    @staticmethod
    def get_all_data(path, sheet_name):

        workbook = load_workbook(path)

        sheet = workbook[sheet_name]

        data = []

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            data.append(row)

        return data